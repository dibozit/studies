from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.drawing.image import Image
from openpyxl import load_workbook


wb= load_workbook('Pie.xlsx') #hali hazırda var olan excel'i çağırdık
ws = wb.active

#tablo eklemek için:
tab= Table(displayName='Table1', ref='A1:B5')
style= TableStyleInfo(name='TableStyleMedium9', showFirstColumn=False, showLastColumn=False,
                                                showRowStripes=True, showColumnStripes=True)
tab.tableStyleInfo = style
ws.add_table(tab)
wb.save('table.xlsx')

#image eklemek için:
img = Image('madecraft.jpg')
img.height = img.height * .5 #boyunu gerçeğinin yarısına indirdim
img.width = img.width * .5 #genişliğini gerçeğinin yarısına indirdim
ws.add_image(img, 'B7')
wb.save('image.xlsx')